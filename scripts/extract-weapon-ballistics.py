#!/usr/bin/env python3
"""Extract the owner-supplied CS2 workbook without executing its formulas/macros.

Requires openpyxl for read-only extraction. The full original stays in private assets.
"""
import argparse
import hashlib
import json
from pathlib import Path
import shutil
import zlib
import openpyxl

parser = argparse.ArgumentParser(description=__doc__)
parser.add_argument('workbook', type=Path)
args = parser.parse_args()
project = Path(__file__).resolve().parents[1]
source = args.workbook
book = openpyxl.load_workbook(source, read_only=True, data_only=False)
cached = openpyxl.load_workbook(source, read_only=True, data_only=True)
mapping = dict(Pistol='deagle', SMG='mp9', Shotgun='xm1014', Assault='ak47', Carbine='scar20', Sniper='awp')
raw = {row[0].value: row for row in book['Raw Values'] if row and row[0].value in mapping.values()}
accuracy_columns = dict(spread='U', crouch='V', stand='W', fire='X', move='Y', jump='Z',
                        jump_initial='AA', jump_apex='AB', land='AC', ladder='AD',
                        recovery_crouch='AE', recovery_crouch_final='AF',
                        recovery_stand='AG', recovery_stand_final='AH', max_speed='I')
alt_columns = dict(spread='AL', crouch='AM', stand='AN', fire='AO', move='AP',
                   jump='AQ', land='AR', ladder='AS', max_speed='AW')
simple_names = dict(Pistol='Desert Eagle', SMG='MP9', Shotgun='XM1014', Assault='AK-47', Carbine='SCAR-20', Sniper='AWP')
simple_rows = {row[0].value: row[0].row for row in book['Simple'] if row and row[0].value in simple_names.values()}
profiles = {}
for key, name in mapping.items():
    row = raw[name]
    number = row[0].row
    values = [float(c.value) for c in row[1:5]]
    armor, damage, modifier, head = values
    assert 0 < modifier <= 1 and damage > 0 and head > 0 and 0 < armor <= 2
    # These are explicit formulas in Analysis, not values inferred from another game.
    analysis = book['Analysis']
    assert '*1.25' in analysis[f'BB{number}'].value
    assert '*0.75' in analysis[f'BC{number}'].value
    samples = []
    for column, distance in [('E', 254), ('I', 1270), ('N', 2540), ('W', 7620)]:
        cell = f'{column}{number}'
        expected = float(cached['Damage @Range'][cell].value)
        assert abs(expected - damage * modifier ** (distance / 1270)) < .00001
        samples.append(dict(cell=cell, distance_cm=distance, damage=expected))
    profiles[key] = dict(row=name, source_range=f'Raw Values!B{number}:E{number}',
                         damage=damage, armor_ratio=armor, range_modifier=modifier,
                         headshot_multiplier=head, stomach_multiplier=1.25,
                         chest_arms_multiplier=1., legs_multiplier=.75, range_samples=samples)
    hip = {prop: float(cached['Raw Values'][f'{column}{number}'].value) for prop, column in accuracy_columns.items()}
    scoped = dict(hip)
    for prop, column in alt_columns.items():
        value = cached['Raw Values'][f'{column}{number}'].value
        if value is not None:
            scoped[prop] = float(value)
    simple_row = simple_rows[simple_names[key]]
    recoil_columns = dict(angle_variance='AI', magnitude='AJ', magnitude_variance='AK')
    recoil_alt_columns = dict(angle_variance='AT', magnitude='AU', magnitude_variance='AV')
    recoil_hip = {prop: float(cached['Raw Values'][f'{column}{number}'].value)
                  for prop, column in recoil_columns.items()}
    recoil_scoped = dict(recoil_hip)
    recoil_alt_cells = {}
    for prop, column in recoil_alt_columns.items():
        value = cached['Raw Values'][f'{column}{number}'].value
        if value is not None:
            recoil_scoped[prop] = float(value)
            recoil_alt_cells[prop] = f'Raw Values!{column}{number}'
    pattern = cached['Simple'][f'AH{simple_row}'].value
    assert pattern in ('Random', 'Set Pattern'), (key, pattern)
    for prop, column in dict(magnitude='AE', angle_variance='AF', magnitude_variance='AG').items():
        assert abs(float(cached['Simple'][f'{column}{simple_row}'].value)-recoil_hip[prop]) < .00001
    profiles[key]['recoil'] = dict(hip=recoil_hip, scoped=recoil_scoped,
        source_cells={prop: f'Raw Values!{column}{number}' for prop, column in recoil_columns.items()},
        alt_source_cells=recoil_alt_cells, pattern=pattern, pattern_cell=f'Simple!AH{simple_row}',
        project_pattern_seed=zlib.crc32(name.encode()) & 0x7fffffff)
    checks = {prop: float(cached['Simple'][f'{column}{simple_row}'].value)
              for prop, column in dict(standing='V', crouching='W', running='X', ladder='Y', jump_apex='Z', landing_333ms='AA').items()}
    assert abs(checks['standing'] - (hip['spread'] + hip['stand'])) < .00001
    assert abs(checks['crouching'] - (hip['spread'] + hip['crouch'])) < .00001
    assert abs(checks['running'] - (hip['spread'] + hip['stand'] + hip['move'])) < .00001
    assert abs(checks['ladder'] - (hip['spread'] + 2*hip['ladder'])) < .00001
    assert abs(checks['jump_apex'] - (hip['spread'] + hip['stand'] + hip['jump'] + hip['jump_apex'])) < .00001
    assert abs(checks['landing_333ms'] - (hip['spread'] + hip['stand'] +
               (hip['jump'] + hip['land'] * 302) * .1**(.333 / hip['recovery_stand']))) < .00001
    bursts = {posture: [float(cached[sheet][f'{openpyxl.utils.get_column_letter(c)}{number}'].value) for c in range(7,17)
                       if cached[sheet][f'{openpyxl.utils.get_column_letter(c)}{number}'].value is not None]
              for posture, sheet in [('standing', 'Firing Inaccuracy(Standing) Raw'), ('crouching', 'Firing Inaccuracy(Crouching) Ra')]}
    tapping = [float(cached['Firing Inaccuracy(Tapping) Raw'][f'{openpyxl.utils.get_column_letter(c)}{number}'].value) for c in range(6,12)]
    tap_interval = float(cached['Firing Inaccuracy(Tapping) Raw'][f'B{number}'].value)
    profiles[key]['accuracy'] = dict(hip=hip, scoped=scoped,
        source_cells={prop: f'Raw Values!{column}{number}' for prop, column in accuracy_columns.items()},
        alt_source_cells={prop: f'Raw Values!{column}{number}' for prop, column in alt_columns.items()
                          if cached['Raw Values'][f'{column}{number}'].value is not None},
        cycle_time=float(cached['Raw Values'][f'F{number}'].value),
        simple_checks=checks, burst_checks=bursts, tapping_checks=tapping, tap_interval=tap_interval,
        standing_full_angle_degrees=float(cached['Analysis'][f'AO{number}'].value),
        standing_full_angle_mode='scoped' if 'AL' in str(book['Analysis'][f'C{number}'].value) else 'hip')
    for posture, expected in bursts.items():
        remaining = 0.
        start = hip['recovery_' + ('stand' if posture == 'standing' else 'crouch')]
        final = hip['recovery_' + ('stand' if posture == 'standing' else 'crouch') + '_final']
        for index, value in enumerate(expected):
            assert abs(remaining - value) < .00001, (key, posture, index, remaining, value)
            recovery = start + (final - start) * min(1, max(0, (index + 1 - 2) / 3))
            remaining = (remaining + hip['fire']) * .1 ** (profiles[key]['accuracy']['cycle_time'] / recovery)
output = dict(source_filename=source.name, sha256=hashlib.sha256(source.read_bytes()).hexdigest(),
              license='Public Domain (License!B4)', range_interval_cm=1270, weapons=profiles)
target = project / 'docs/references/cs2-weapon-ballistics-2026-03-18.json'
target.write_text(json.dumps(output, indent=2) + '\n')
archive = project / 'Assets/weapons/SourceAssets/References/CS2-Weapon-Spreadsheet-2026-03-18.xlsx'
archive.parent.mkdir(parents=True, exist_ok=True)
if archive.exists():
    assert archive.read_bytes() == source.read_bytes(), 'Different source archive already exists'
else:
    shutil.copy2(source, archive)
print(target)
print('Extracted six profiles; checked 24 distance samples against cached workbook results.')
